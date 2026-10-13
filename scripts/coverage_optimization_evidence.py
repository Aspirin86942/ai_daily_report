"""Source-backed, credential-free evidence for coverage optimization."""

from __future__ import annotations

import argparse
from io import BytesIO
from datetime import UTC, date, datetime, timedelta
from hashlib import sha256
import json
import os
from pathlib import Path
import subprocess
import sys
import zipfile
from time import perf_counter
from types import SimpleNamespace
from urllib.parse import urlsplit

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from src.core.config import Config
from src.models.report_settings import ModelRequestSettings
from src.services.native_scanner import NativeScanner, ScanRequest
from src.services.report_runner.context import RunContext
from src.services.report_runner.model_plan import plan_evidence


def effective_snapshot(runtime: Config) -> dict:
    """Export approved settings only; never export YAML or environment values."""
    scanner = runtime.scanner_settings()
    public_scanner = {}
    for key, value in scanner.items():
        if key in {"excluded_dirs", "office_worker_path", "index_db_path"}:
            public_scanner[key] = {"configured": bool(value), "count": len(value) if isinstance(value, list) else 1}
        else:
            public_scanner[key] = value
    llm = runtime.llm_config
    return {
        "version": 1,
        "python": sys.version.split()[0],
        "root_alias": runtime.report_settings.root_alias,
        "work_root_fingerprint": sha256(str(runtime.work_dir.resolve()).encode()).hexdigest(),
        "report": runtime.report_settings.model_dump(exclude={"notes_dir"}),
        "scanner": public_scanner,
        "llm": {key: llm[key] for key in ("model_id", "temperature", "max_tokens", "max_retries")},
        "gateway_host": urlsplit(llm["base_url"]).hostname,
        "provider": runtime.llm_provider,
        "model_requests": runtime.model_request_settings.model_dump(),
        "recovery": runtime.recovery_settings.model_dump(),
        "schedule": runtime.schedule_settings.model_dump(),
        "retention": runtime.retention_settings.model_dump(),
        "environment_overrides": sorted(key for key in os.environ if key.startswith("DAILY_REPORT_")),
        "credential_present": "DEEPSEEK_API_KEY" in os.environ,
    }


def pressure_corpus(root: Path, target: date) -> list[dict]:
    """Deterministic synthetic facts; refuse to overwrite any existing input."""
    root.mkdir(parents=True, exist_ok=False)
    modified = datetime.combine(target, datetime.min.time(), UTC) + timedelta(hours=1)
    facts = []
    for index in range(9):
        name = f"000-large-{index:02d}.md"
        fact = f"材料记载项目 L{index:02d} 的样本金额为 {10000 + index} 元，尚待业务确认。"
        text = f"# 项目 L{index:02d}\n\n" + (fact + "\n") * 240
        path = root / name
        path.write_text(text, encoding="utf-8", newline="\n")
        os.utime(path, (modified.timestamp(), modified.timestamp()))
        facts.append({"path": name, "fact": fact, "sha256": sha256(path.read_bytes()).hexdigest()})
    for index in range(12):
        name = f"zzz-small-{index:02d}.md"
        fact = f"材料记载项目 S{index:02d} 的合同样本金额为 {20000 + index} 元，单位为人民币元。"
        path = root / name
        path.write_text(f"# 项目 S{index:02d}\n\n{fact}\n", encoding="utf-8", newline="\n")
        os.utime(path, (modified.timestamp(), modified.timestamp()))
        facts.append({"path": name, "fact": fact, "sha256": sha256(path.read_bytes()).hexdigest()})
    return facts


def mixed_corpus(root: Path, target: date) -> list[dict]:
    """Freeze synthetic numeric facts and boundary cases before measuring new parsers."""
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from reportlab.pdfgen import canvas

    root.mkdir(parents=True, exist_ok=False)
    facts = []
    modified = datetime.combine(target, datetime.min.time(), UTC) + timedelta(hours=1)

    def record(name, expected, **limits):
        path = root / name
        os.utime(path, (modified.timestamp(), modified.timestamp()))
        facts.append({"path": name, "sha256": sha256(path.read_bytes()).hexdigest(),
                      "expected_facts": expected, "boundaries": limits})

    def office_bytes(document, name):
        stream = BytesIO()
        document.save(stream)
        result = BytesIO()
        with zipfile.ZipFile(stream) as archive, zipfile.ZipFile(result, "w", compression=zipfile.ZIP_DEFLATED) as frozen:
            for member in sorted(archive.namelist()):
                info = zipfile.ZipInfo(member, (2026, 9, 22, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                frozen.writestr(info, archive.read(member))
        (root / name).write_bytes(result.getvalue())

    doc = Document()
    doc.add_heading("合同验收样本", 0)
    doc.add_paragraph("方案记载验收日期为2026-09-22，计划金额为人民币1234.56元。")
    table = doc.add_table(rows=1, cols=4)
    for cell, value in zip(table.rows[0].cells, ("项目", "数量", "单位", "金额（人民币元）")):
        cell.text = value
    for values in (("A", "2", "件", "400.00"), ("B", "3", "件", "834.56"), ("局部小计", "", "元", "1234.56")):
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    office_bytes(doc, "contract.docx")
    record("contract.docx", ["2026-09-22", "人民币1234.56元", "局部小计", "834.56"], paragraphs=2, table_rows=4)
    office_bytes(Document(), "empty.docx")
    record("empty.docx", [], valid_empty_document=True, placeholder_must_not_be_evidence=True)
    (root / "zero.docx").write_bytes(b"")
    record("zero.docx", [], expected_failure="empty_container")

    book = Workbook()
    sheet = book.active
    sheet.title = "采购样本"
    sheet.append(["币种及单位：人民币元"])
    sheet.append(["编号", "数量（件）", "金额（元）"])
    for number in range(1, 55):
        sheet.append([f"ITEM-{number:02d}", number, number * 100])
    sheet.append(["全表合计", 1485, 148500])
    second = book.create_sheet("局部汇总")
    second.append(["项目", "局部小计（元）"])
    second.append(["样本", 1234.56])
    office_bytes(book, "amounts.xlsx")
    record("amounts.xlsx", ["人民币元", "ITEM-01", "1234.56"], sheets=2, rows_first_sheet=57,
           outside_default_prefix={"sheet": "采购样本", "row": 57, "label": "全表合计", "amount": 148500})
    (root / "modern_mislabeled.xls").write_bytes((root / "amounts.xlsx").read_bytes())
    record("modern_mislabeled.xls", ["ITEM-01", "1234.56"], actual_format="xlsx", declared_extension=".xls")
    (root / "legacy.xls").write_bytes((ROOT / "tests/fixtures/worker_documents/legacy_sample.xls").read_bytes())
    record("legacy.xls", ["Legacy XLS worker content"], actual_format="xls", requires_legacy_enabled=True)

    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = "验收计划"
    slide.placeholders[1].text = "计划于2026-09-22验收，预算1234.56元。"
    slide.notes_slide.notes_text_frame.text = "备注：这是计划日期，尚未证明实际完成。"
    office_bytes(slides, "plan.pptx")
    record("plan.pptx", ["2026-09-22", "1234.56", "尚未证明实际完成"], slides=1, notes_required=True)

    for name, pages, text_page in (("later_text.pdf", 8, 8), ("no_text.pdf", 8, None), ("early_text.pdf", 2, 2)):
        document = canvas.Canvas(str(root / name), invariant=1)
        for page in range(1, pages + 1):
            if page == text_page:
                document.drawString(72, 720, "SYNTH-AMOUNT 1234.56 CNY; planned 2026-09-22")
            document.showPage()
        document.save()
        record(name, ["1234.56", "2026-09-22"] if text_page else [], pages=pages, text_page=text_page,
               text_outside_initial_window=bool(text_page and text_page > 5))
    (root / "damaged.pdf").write_bytes(b"%PDF-1.7\nBROKEN-SYNTHETIC-CONTAINER")
    record("damaged.pdf", [], expected_failure="corrupt_pdf")

    csv_text = '编号,金额（人民币元）,备注\nSYNTH-1,1234.56,"计划2026-09-22\n尚未执行"\nSYNTH-2,100.00,𠀀\n'
    for name, encoding in (("utf8.csv", "utf-8"), ("bom.csv", "utf-8-sig"), ("gb18030.csv", "gb18030")):
        (root / name).write_bytes(csv_text.encode(encoding))
        record(name, ["1234.56", "2026-09-22", "𠀀"], encoding=encoding, logical_records=3,
               gb18030_requires_explicit_allowlist=encoding == "gb18030")
    return facts


def synthetic_baseline(directory: Path, runtime_root: Path, target: date) -> dict:
    corpus = directory / "work"
    truth = pressure_corpus(corpus, target)
    profile = {"allowed_extensions": [".md"], "total_max_chars": 50000,
               "text_max_chars": 6000, "total_deadline_ms": 10000}
    scanner = NativeScanner(SimpleNamespace(
        work_dir=corpus, index_db_path=directory / "scan_index_v3.sqlite3",
        office_worker_path=runtime_root / "rust/target/release/ai-daily-office-parser.exe",
        scanner_settings=lambda: profile,
    ))
    context = RunContext.create(now=datetime.combine(target + timedelta(days=1), datetime.min.time(), UTC),
                                target_date=target)
    request = ScanRequest("daily", target, target, window_start_utc=context.window_start_utc,
                          window_end_utc=context.window_end_utc)
    results = []
    for mode in ("cold", "warm", "snapshot"):
        started = perf_counter()
        active_request = request if mode == "cold" else ScanRequest("daily", target - timedelta(days=1), target)
        scan = scanner.build_context(active_request)
        elapsed = round((perf_counter() - started) * 1000, 3)
        plan = plan_evidence(context, scan.sources, user_input="", system_prompt="仅返回有来源的 JSON 摘要。",
            llm_config={"base_url": "https://synthetic.invalid", "provider": "deepseek",
                        "model_id": "synthetic", "temperature": 0.2, "max_tokens": 8192},
            settings=ModelRequestSettings(context_window_tokens=1000000, verified_max_output_tokens=8192),
            scanner_profile=profile)
        results.append({"mode": mode, "status": scan.envelope.status, "elapsed_ms": elapsed,
            "registered": len(scan.sources), "nonempty": sum(bool(s.excerpt.strip()) for s in scan.sources),
            "excerpt_chars": sum(len(s.excerpt) for s in scan.sources),
            "planned_covered": sum(bool(s.model_chars) for s in plan.sources),
            "facts_found": sum(any(case["fact"] in source.excerpt for source in scan.sources) for case in truth),
            "planned_batches": len(plan.batches), "input_fingerprint": plan.input_fingerprint,
            "execution_metrics": scan.evidence.execution_metrics.model_dump(mode="json") if scan.evidence else None,
            "sources": [{"path": s.relative_path, "parse_status": s.parse_status,
                         "excerpt_chars": len(s.excerpt), "excerpt_sha256": s.excerpt_sha256,
                         "content_sha256": s.content_sha256, "reason": s.reason,
                         "coverage_limits": s.coverage_limits} for s in scan.sources]})
    return {"profile": profile, "target_date": str(target), "truth": truth,
            "model_calls": 0, "results": results}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--runtime-root", type=Path, required=True)
    parser.add_argument("--worker-root", type=Path, default=ROOT)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--target-date", type=date.fromisoformat, default=date(2026, 10, 11))
    parser.add_argument("--config-only", action="store_true")
    parser.add_argument("--mixed-corpus-only", action="store_true")
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=False)
    if args.mixed_corpus_only:
        truth = mixed_corpus(args.output / "work", args.target_date)
        path = args.output / "truth.json"
        path.write_text(json.dumps({"version": 1, "target_date": str(args.target_date), "cases": truth},
                                  ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
        print(json.dumps({"truth": str(path), "cases": len(truth), "model_calls": 0}, ensure_ascii=False))
        return 0
    runtime = Config()
    runtime._project_root = args.runtime_root.resolve()
    runtime._settings = Config._build_settings(args.runtime_root / "config")
    evidence = {"version": 1, "head": subprocess.check_output(
        ["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip(),
        "config": effective_snapshot(runtime)}
    if not args.config_only:
        evidence["synthetic"] = synthetic_baseline(args.output / "synthetic", args.worker_root, args.target_date)
    path = args.output / "evidence.json"
    path.write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    print(json.dumps({"path": str(path), "model_calls": 0,
                      "synthetic": [{key: row[key] for key in (
                          "mode", "status", "registered", "nonempty", "planned_covered", "elapsed_ms")}
                          for row in evidence.get("synthetic", {}).get("results", [])]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
