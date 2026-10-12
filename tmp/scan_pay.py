import openpyxl

# 1) 付款申请单 rows in 刘智流程统计.xlsx -> node names
f = r'd:\01- 工作\2026年\32- 流程梳理\刘智流程统计\刘智流程统计.xlsx'
wb = openpyxl.load_workbook(f, data_only=True, read_only=True)
ws = wb['原始数据']
print('==== 付款申请单 rows (刘智流程统计)')
n = 0
for row in ws.iter_rows(values_only=True):
    joined = ' | '.join('' if v is None else str(v) for v in row)
    if '付款申请单' in joined:
        n += 1
        if n <= 40:
            print('  ', joined[:220])
print('  total 付款申请单 rows:', n)
wb.close()

# 2) unique node names for 付款申请单 flows
wb = openpyxl.load_workbook(f, data_only=True, read_only=True)
ws = wb['原始数据']
nodes = set()
header = None
for i, row in enumerate(ws.iter_rows(values_only=True)):
    vals = ['' if v is None else str(v) for v in row]
    if i == 0:
        header = vals
        continue
    joined = ' | '.join(vals)
    if '付款申请单' in joined and len(vals) > 1:
        nodes.add(vals[1])
print('==== unique node names for 付款申请单:')
for x in sorted(nodes):
    print('  ', x)
wb.close()

# 3) 韩老师 payment export header
f2 = r'd:\01- 工作\2026年\32- 流程梳理\韩老师\付款申请单_2026081017561720_101399.xlsx'
wb2 = openpyxl.load_workbook(f2, data_only=True, read_only=True)
print('==== 韩老师\付款申请单 sheets:', wb2.sheetnames)
for ws in wb2.worksheets:
    print('-- sheet', ws.title)
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        vals = ['' if v is None else str(v) for v in row]
        if any(vals):
            print('   ', ' | '.join(vals)[:300])
        if i > 5:
            break
wb2.close()
