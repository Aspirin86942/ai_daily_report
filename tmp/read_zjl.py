import openpyxl

path = r'd:\01- 工作\2026年\32- 流程梳理\总经理\流程统计\流程统计-总经理.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
print('SHEETS:', wb.sheetnames)
for ws in wb.worksheets:
    print('=' * 100)
    print('SHEET:', ws.title, 'dims:', ws.dimensions)
    for row in ws.iter_rows(values_only=True):
        vals = ['' if v is None else str(v).replace('\n', '⏎') for v in row]
        if any(vals):
            print(' | '.join(vals))
